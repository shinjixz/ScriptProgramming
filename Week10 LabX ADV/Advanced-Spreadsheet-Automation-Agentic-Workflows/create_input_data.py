# agentic-spreadsheet-processor/create_input_data.py
# Generates data/input_data.xlsx with the sample sales data used in this lab.
import os

import openpyxl
from openpyxl.styles import Font

SALES = [
    ("Product Name", "Quantity", "Unit Price"),
    ("Laptop", 2, 1200.50),
    ("Mouse", 5, 25.00),
    ("Keyboard", 3, 75.99),
    ("Monitor", 1, 300.00),
    ("Projector", 1, 850.00),
    ("Webcam", 10, 45.00),
    ("Headset", 4, 99.99),
    ("SSD", 2, 180.00),
    ("Router", 1, 70.00),
]


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(base, "data", "input_data.xlsx")
    os.makedirs(os.path.dirname(path), exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SalesData"
    for row in SALES:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "#,##0.00"
    for col in "ABC":
        ws.column_dimensions[col].width = 15
    wb.save(path)
    print(f"Created {path}")


if __name__ == "__main__":
    main()
