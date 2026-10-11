# agentic-spreadsheet-processor/src/spreadsheet_agent.py
import datetime
import json
import os

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule, Rule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries

from .utils import setup_logging

logger = setup_logging(__name__)


class SpreadsheetAgent:
    """
    An agent that performs a series of operations on an Excel workbook
    based on a provided configuration.
    """

    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()
        logger.info("Spreadsheet Agent initialized.")

    # ------------------------------------------------------------------ #
    # Helpers
    # ------------------------------------------------------------------ #
    def _log_audit(self, status, message, task_type="N/A", details=None):
        """Adds an entry to the audit log."""
        log_entry = {
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {},
        }
        self.audit_log.append(log_entry)
        if status in ("ERROR", "CRITICAL"):
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type} | Details: {details}")
        elif status == "WARNING":
            logger.warning(f"AUDIT LOG - {status}: {message} | Task: {task_type}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        """Loads the input workbook."""
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook")
            return True
        except FileNotFoundError:
            self._log_audit("ERROR", f"Input file not found: {input_file}", "load_workbook")
            return False
        except Exception as e:
            self._log_audit("ERROR", f"Failed to load workbook '{input_file}': {e}", "load_workbook")
            return False

    def _get_sheet(self, sheet_name, create=True):
        """Safely gets a worksheet from the workbook (optionally creating it)."""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        if not create:
            return None
        self._log_audit("WARNING", f"Sheet '{sheet_name}' not found. Creating new sheet.", "get_sheet")
        return self.workbook.create_sheet(sheet_name)

    @staticmethod
    def _last_data_row(sheet, column="A"):
        """Returns the last row that has a value in the given column."""
        for row in range(sheet.max_row, 0, -1):
            if sheet[f"{column}{row}"].value not in (None, ""):
                return row
        return 1

    # ------------------------------------------------------------------ #
    # Task handlers
    # ------------------------------------------------------------------ #
    def _process_copy_data_task(self, task):
        """Copies data from a source sheet to a destination sheet."""
        src_sheet_name = task.get("source_sheet")
        dest_sheet_name = task.get("dest_sheet")
        if not src_sheet_name or not dest_sheet_name:
            self._log_audit("ERROR", "Missing source_sheet or dest_sheet for copy_data task.", "copy_data", task)
            return

        src_sheet = self._get_sheet(src_sheet_name, create=False)
        if src_sheet is None:
            self._log_audit("ERROR", f"Source sheet '{src_sheet_name}' not available for copying.", "copy_data", task)
            return
        dest_sheet = self._get_sheet(dest_sheet_name)

        try:
            rows = 0
            for row in src_sheet.iter_rows():
                for cell in row:
                    dest_sheet[cell.coordinate].value = cell.value
                    if cell.has_style:
                        dest_sheet[cell.coordinate].number_format = cell.number_format
                rows += 1
            # Make header bold and widen columns a little for readability
            for cell in dest_sheet[1]:
                cell.font = Font(bold=True)
            for col_idx in range(1, src_sheet.max_column + 1):
                letter = get_column_letter(col_idx)
                dest_sheet.column_dimensions[letter].width = 15
            self._log_audit("SUCCESS", f"Data copied from '{src_sheet_name}' to '{dest_sheet_name}' ({rows} rows).",
                            "copy_data")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to copy data: {e}", "copy_data", task)

    def _process_calculate_column_task(self, task):
        """Calculates a new column based on a formula or existing columns."""
        sheet_name = task.get("sheet")
        target_column_letter = task.get("target_column")
        start_row = task.get("start_row", 2)  # Assume header in row 1
        formula = task.get("formula")  # e.g., "=B{row}*C{row}"
        if not sheet_name or not target_column_letter or not formula:
            self._log_audit("ERROR", "Missing required fields for calculate_column task.", "calculate_column", task)
            return

        sheet = self._get_sheet(sheet_name)
        header_name = task.get("header", "Calculated Column")
        sheet[f"{target_column_letter}1"].value = header_name
        sheet[f"{target_column_letter}1"].font = Font(bold=True)
        sheet.column_dimensions[target_column_letter].width = 15

        last_row = self._last_data_row(sheet, "A")
        errors = 0
        for row_num in range(start_row, last_row + 1):
            cell_coord = f"{target_column_letter}{row_num}"
            try:
                sheet[cell_coord] = formula.replace("{row}", str(row_num))
                if task.get("number_format"):
                    sheet[cell_coord].number_format = task["number_format"]
            except Exception as e:
                errors += 1
                self._log_audit("ERROR", f"Failed to apply formula to {cell_coord}: {e}", "calculate_column", task)

        status = "SUCCESS" if errors == 0 else "WARNING"
        self._log_audit(status,
                        f"Column '{target_column_letter}' ({header_name}) calculated in '{sheet_name}' "
                        f"for rows {start_row}-{last_row}.",
                        "calculate_column")

    def _process_conditional_format_task(self, task):
        """Applies conditional formatting to a range."""
        sheet_name = task.get("sheet")
        data_range = task.get("range")
        rule_type = task.get("rule_type")
        if not sheet_name or not data_range or not rule_type:
            self._log_audit("ERROR", "Missing required fields for conditional_format task.", "conditional_format", task)
            return

        sheet = self._get_sheet(sheet_name)
        try:
            if rule_type == "color_scale":
                rule = ColorScaleRule(start_type='min', start_color=task.get("min_color", "FF0000"),
                                      mid_type='percentile', mid_value=50, mid_color=task.get("mid_color", "FFFF00"),
                                      end_type='max', end_color=task.get("max_color", "00FF00"))
            elif rule_type == "data_bar":
                rule = DataBarRule(start_type='min', end_type='max', color=task.get("color", "0070C0"))
            elif rule_type == "icon_set":
                rule = IconSetRule(task.get("icon_set", '3TrafficLights1'),
                                   task.get("value_type", 'percent'),
                                   task.get("values", [0, 33, 67]))
            elif rule_type == "expression":
                formula = task.get("formula", "").lstrip("=")
                dxf = DifferentialStyle(
                    font=Font(color=task.get("font_color", "9C0006")),
                    fill=PatternFill(bgColor=task.get("fill_color", "FFC7CE"), fill_type="solid"),
                )
                rule = Rule(type="expression", formula=[formula], dxf=dxf)
            else:
                self._log_audit("ERROR", f"Unsupported conditional format rule_type: {rule_type}",
                                "conditional_format", task)
                return

            sheet.conditional_formatting.add(data_range, rule)
            self._log_audit("SUCCESS", f"Conditional formatting '{rule_type}' applied to '{data_range}' in "
                                       f"'{sheet_name}'.", "conditional_format")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to apply conditional formatting to '{data_range}': {e}",
                            "conditional_format", task)

    def _process_chart_task(self, task):
        """Creates a chart based on specified data."""
        sheet_name = task.get("sheet")
        chart_type = task.get("chart_type")
        data_range = task.get("data_range")  # e.g., "D1:D5" (header included)
        category_range = task.get("category_range")  # e.g., "A2:A5" for labels
        title = task.get("title", "Chart Title")
        top_left_cell = task.get("top_left_cell", "G2")
        if not sheet_name or not chart_type or not data_range:
            self._log_audit("ERROR", "Missing required fields for chart task.", "chart", task)
            return

        sheet = self._get_sheet(sheet_name)
        try:
            charts = {"bar": BarChart, "line": LineChart, "pie": PieChart}
            if chart_type not in charts:
                self._log_audit("ERROR", f"Unsupported chart type: {chart_type}", "chart", task)
                return
            chart = charts[chart_type]()

            min_col, min_row, max_col, max_row = range_boundaries(data_range)
            data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True)  # first row = series title

            if category_range:
                c_min_col, c_min_row, c_max_col, c_max_row = range_boundaries(category_range)
                cats = Reference(sheet, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
                chart.set_categories(cats)

            chart.title = title
            if chart_type != "pie":
                chart.x_axis.title = task.get("x_axis_title", "")
                chart.y_axis.title = task.get("y_axis_title", "")
                # openpyxl >= 3.1 hides axes by default; turn them back on
                chart.x_axis.delete = False
                chart.y_axis.delete = False
            chart.width = task.get("width", 16)
            chart.height = task.get("height", 8)

            sheet.add_chart(chart, top_left_cell)
            self._log_audit("SUCCESS", f"Chart '{title}' of type '{chart_type}' added to '{sheet_name}' "
                                       f"at {top_left_cell}.", "chart")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to create chart '{title}': {e}", "chart", task)

    def _process_summary_task(self, task):
        """Creates a summary sheet whose cells are Excel formulas over another sheet."""
        sheet_name = task.get("sheet")
        metrics = task.get("metrics", [])
        if not metrics:
            self._log_audit("ERROR", "Summary task requires a non-empty 'metrics' list.", "summary", task)
            return

        sheet = self._get_sheet(sheet_name)
        try:
            sheet["A1"] = task.get("title", "Summary")
            sheet["A1"].font = Font(bold=True, size=14)
            sheet["A3"], sheet["B3"] = "Metric", "Value"
            for c in (sheet["A3"], sheet["B3"]):
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
            for i, metric in enumerate(metrics, start=4):
                sheet[f"A{i}"] = metric["label"]
                sheet[f"B{i}"] = metric["formula"]
                if metric.get("number_format"):
                    sheet[f"B{i}"].number_format = metric["number_format"]
                sheet[f"B{i}"].alignment = Alignment(horizontal="right")
            sheet.column_dimensions["A"].width = 28
            sheet.column_dimensions["B"].width = 18
            self._log_audit("SUCCESS", f"Summary sheet '{sheet_name}' created with {len(metrics)} metrics.", "summary")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to build summary: {e}", "summary", task)

    # ------------------------------------------------------------------ #
    # Reporting
    # ------------------------------------------------------------------ #
    def _add_audit_log_sheet(self):
        """Adds a sheet to the workbook containing the audit log."""
        if not self.audit_log:
            return
        if "Audit Log" in self.workbook.sheetnames:
            del self.workbook["Audit Log"]
        audit_sheet = self.workbook.create_sheet("Audit Log", 0)  # Create as first sheet
        audit_sheet.append(["Timestamp", "Status", "Task Type", "Message", "Details"])

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
        for col_num in range(1, 6):
            cell = audit_sheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill

        status_colors = {"SUCCESS": "C6EFCE", "INFO": "DDEBF7", "WARNING": "FFEB9C",
                         "ERROR": "FFC7CE", "CRITICAL": "FF8080"}
        for entry in self.audit_log:
            audit_sheet.append([
                entry["timestamp"],
                entry["status"],
                entry["task_type"],
                entry["message"],
                json.dumps(entry["details"], ensure_ascii=False),
            ])
            color = status_colors.get(entry["status"])
            if color:
                audit_sheet.cell(row=audit_sheet.max_row, column=2).fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid")

        # Auto-adjust column widths
        for col_num in range(1, 6):
            column_letter = get_column_letter(col_num)
            max_length = max((len(str(c.value)) for c in audit_sheet[column_letter] if c.value), default=0)
            audit_sheet.column_dimensions[column_letter].width = min(max_length + 2, 80)
        audit_sheet.freeze_panes = "A2"
        logger.info("Audit log sheet created.")

    # ------------------------------------------------------------------ #
    # Main loop
    # ------------------------------------------------------------------ #
    def run(self):
        """Executes the defined tasks from the configuration."""
        if not self._load_workbook():
            self._log_audit("CRITICAL", "Agent could not start due to workbook loading failure.", "run")
            return False

        handlers = {
            "copy_data": self._process_copy_data_task,
            "calculate_column": self._process_calculate_column_task,
            "conditional_format": self._process_conditional_format_task,
            "chart": self._process_chart_task,
            "summary": self._process_summary_task,
        }

        try:
            for i, task in enumerate(self.config["tasks"], start=1):
                task_type = task.get("type")
                logger.info(f"Executing Task {i}: Type='{task_type}' Sheet='{task.get('sheet', task.get('dest_sheet'))}'")
                handler = handlers.get(task_type)
                if handler:
                    handler(task)
                else:
                    self._log_audit("WARNING", f"Unsupported task type encountered: {task_type}",
                                    "task_execution", task)

            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            n_err = sum(1 for e in self.audit_log if e["status"] in ("ERROR", "CRITICAL"))
            self._log_audit("INFO", f"Agent executed {len(self.config['tasks'])} tasks in {duration:.2f} seconds "
                                    f"({n_err} errors).", "run_summary")

            self._add_audit_log_sheet()

            output_file = self.config["output_file"]
            os.makedirs(os.path.dirname(output_file) or ".", exist_ok=True)
            self.workbook.save(output_file)
            logger.info(f"Final workbook saved to: {output_file}")
            return True
        except Exception as e:
            self._log_audit("CRITICAL", f"An unhandled error occurred during agent execution: {e}", "run")
            logger.critical(f"An unhandled error occurred: {e}", exc_info=True)
            return False
